"""Bounded, read-only parsers. Locations refer to extracted text, never guessed bytes."""
from __future__ import annotations

import hashlib
import io
import zipfile
from pathlib import Path

VERSION = 'source-segments/1'
MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 4_000_000
MAX_SEGMENTS = 4000
SUPPORTED = {'.txt', '.md', '.bdf', '.nas', '.pdf', '.docx'}


def sha(text: str) -> str:
    return hashlib.sha256(text.encode('utf-8')).hexdigest()


def parse(raw: bytes, name: str) -> dict:
    """Return every parsed segment or fail explicitly; scanned pages are reported."""
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError('支持 TXT、MD、BDF、NAS、文字型 PDF 和 DOCX；不支持旧版 DOC 或 CAD 语义解析')
    if not raw or len(raw) > MAX_BYTES:
        raise ValueError('文件为空或超过 20 MiB；请按资料边界分批导入')
    parents, missing = [], []
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError('加密 PDF 请先在本机解密')
        if len(reader.pages) > 2000:
            raise ValueError('PDF 超过 2000 页；请按章节分批导入')
        for page_no, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if not text.strip():
                missing.append({'page': page_no, 'reason': '无可提取文字，可能需要 OCR'})
            else:
                parents.append((page_no, f'第 {page_no} 页（提取文本）', text))
    elif ext == '.docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError('DOCX 解压规模超过 80 MiB')
            if len(archive.infolist()) > 10000:
                raise ValueError('DOCX 内部条目过多')
        document = Document(io.BytesIO(raw))
        for i, p in enumerate(document.paragraphs, 1):
            if p.text.strip():
                parents.append((0, f'正文段落 {i}（不声明页码）', p.text))
        for ti, table in enumerate(document.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                text = '\t'.join(cell.text for cell in row.cells)
                if text.strip():
                    parents.append((0, f'表 {ti} 行 {ri}（不声明页码）', text))
        missing.append({'reason': 'DOCX 首版仅解析正文段落及顶层表格；页眉页脚、文本框和图片未提取'})
    else:
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = raw.decode('gb18030')
            except UnicodeDecodeError as exc:
                raise ValueError('文本编码不是 UTF-8 / GB18030；没有使用有损替换') from exc
        if '\x00' in text:
            raise ValueError('文件包含 NUL 字节，不按普通文本解析')
        # BDF INCLUDE must not cause an extractor to read outside the registered asset.
        if ext in {'.bdf', '.nas'}:
            import re
            if re.search(r'^\s*INCLUDE\b', text, re.I | re.M):
                missing.append({'reason': '发现 INCLUDE：本次仅处理当前原件，不读取外部包含文件'})
        parents.append((0, '提取正文（字符位置，不是字节位置）', text))
    if sum(len(p[2]) for p in parents) > MAX_TEXT:
        raise ValueError('提取文本超过 400 万字符；请分批处理，未静默截断')
    segments = []
    for page, locator, text in parents:
        parent_hash = sha(text)
        for start in range(0, len(text), 2800):
            end = min(len(text), start + 3000)
            part = text[start:end]
            if not part.strip():
                continue
            segments.append({'page': page, 'locator': locator, 'start': start, 'end': end,
                             'text': part, 'text_hash': sha(part), 'parent_hash': parent_hash})
            if end == len(text):
                break
    if not segments:
        raise ValueError('没有可提取正文；扫描资料需要单独 OCR，不会标记为导入完成')
    if len(segments) > MAX_SEGMENTS:
        raise ValueError('片段超过 4000 个；请分批处理，未静默截断')
    return {'parser': VERSION, 'segments': segments, 'gaps': missing,
            'coverage': 'partial' if missing else 'parsed_text_complete',
            'engineering_parse': False}
